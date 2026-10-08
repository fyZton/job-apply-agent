"""Application log: the user's Excel sheet plus a local history so offers are never processed twice."""
import datetime as dt
import json
import shutil
from copy import copy
from pathlib import Path

import openpyxl

from jobagent.sites import offer_key

COLUMNS = ["Date", "Company", "Title", "Lane", "Source", "Link", "Fit (1-10)",
           "CV used", "Status", "Next step", "Notes"]
SHEET = "Applications"
LINK_COL = COLUMNS.index("Link")


FORMULA_STARTS = ("=", "+", "-", "@", chr(9), chr(13))


def _safe_cell(value):
    """Prefixes text a spreadsheet would read as a formula (= + - @, tab, CR) with a single quote."""
    return "'" + value if isinstance(value, str) and value.startswith(FORMULA_STARTS) else value


def today():
    return dt.date.today().isoformat()


class Tracker:
    def __init__(self, excel, data_dir, log):
        self.excel = Path(excel)
        self.data = Path(data_dir)
        self.data.mkdir(parents=True, exist_ok=True)
        self.log = log
        self.history_path = self.data / "history.json"
        self.pending_path = self.data / "pending_rows.json"
        self.history = self._read(self.history_path, {"seen": {}, "count": {}})
        self.pending = self._read(self.pending_path, [])
        self._backed_up = False
        self._warned_open = False
        if not self.excel.exists():
            self._create_workbook()
        self.excel_keys = self._excel_keys()
        if self.pending:
            self._flush()

    # ---------- history ----------
    @staticmethod
    def _read(path, default):
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError):
            return default

    def _write(self, path, data):
        path.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")

    def _excel_keys(self):
        wb = openpyxl.load_workbook(self.excel, read_only=True)
        keys = {offer_key(r[LINK_COL]) for r in wb[SHEET].iter_rows(min_row=2, values_only=True)
                if r and len(r) > LINK_COL and r[LINK_COL]}
        wb.close()
        return keys

    def seen(self, k):
        return k in self.excel_keys or k in self.history["seen"]

    def mark(self, k, result, **extra):
        self.history["seen"][k] = {"date": today(), "result": result, **extra}
        self._write(self.history_path, self.history)

    def sent_today(self, site):
        return self.history["count"].get(today(), {}).get(site, 0)

    def add_sent(self, site):
        day = self.history["count"].setdefault(today(), {})
        day[site] = day.get(site, 0) + 1
        self._write(self.history_path, self.history)

    # ---------- Excel ----------
    def _create_workbook(self):
        self.excel.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append(COLUMNS)
        wb.save(self.excel)

    def add_row(self, row):
        """row: dict keyed by COLUMNS (Date is filled in here)."""
        self.excel_keys.add(offer_key(row.get("Link")))
        self.pending.append({**{k: _safe_cell(v) for k, v in row.items()}, "Date": today()})
        self._flush()

    def _backup(self):
        if self._backed_up:
            return
        folder = self.data / "backups"
        folder.mkdir(exist_ok=True)
        target = folder / f"{self.excel.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
        shutil.copy2(self.excel, target)
        for old in sorted(folder.glob(f"{self.excel.stem}_*.xlsx"))[:-10]:
            old.unlink()
        self._backed_up = True

    @staticmethod
    def _first_empty(ws):
        """First gap from the top, to stay inside the range of formulas and drop-down lists."""
        n = 2
        while any(ws.cell(n, c).value not in (None, "") for c in range(1, len(COLUMNS) + 1)):
            n += 1
        return n

    @staticmethod
    def _put(ws, n, row):
        for c, name in enumerate(COLUMNS, start=1):
            cell = ws.cell(n, c)
            value = row.get(name, "")
            if name == "Date":
                value = dt.datetime.fromisoformat(value)
            cell.value = value
            if cell.data_type == "f":  # web text starting with "=" must stay text, never a formula
                cell.data_type = "s"
            if n > 2:
                above = ws.cell(n - 1, c)  # the row above always has data (n is the first gap)
                if above.has_style:
                    cell.font = copy(above.font)
                    cell.fill = copy(above.fill)
                    cell.border = copy(above.border)
                    cell.alignment = copy(above.alignment)
                    cell.number_format = above.number_format

    def _flush(self):
        """Writes pending rows to Excel. If the file is open (locked), keeps them for the next run."""
        if not self.pending:
            return True
        try:
            self._backup()
            wb = openpyxl.load_workbook(self.excel)
            ws = wb[SHEET]
            n = self._first_empty(ws)
            for row in self.pending:
                self._put(ws, n, row)
                n += 1
            wb.save(self.excel)
        except PermissionError:
            self._write(self.pending_path, self.pending)
            if not self._warned_open:
                self.log("WARNING: the Excel file is open. Rows are kept and written once it is closed.")
                self._warned_open = True
            return False
        self.pending = []
        self._write(self.pending_path, self.pending)
        return True

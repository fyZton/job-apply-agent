"""Draft profile from an existing CV (`python -m jobagent --setup --from-cv cv.pdf`).

The draft is only a starting point for the interview: nothing is written until the user confirms there.
CV text is treated as untrusted. Contact details come from regexes; facts come from the LLM, and a fact is
kept only if the quote the model gives as its source really appears in the CV and mentions the fact.
"""
import re
from collections import Counter
from pathlib import Path

from jobagent import llm
from jobagent.profile import ID_PREFIX, SCHEMA_VERSION, fact_errors, slug
from jobagent.safety import looks_injected

MISSING = "Reading {} files needs the cv extra: pip install 'jobagent[cv]'"
EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE = re.compile(r"\+?\d[\d\s().-]{7,}\d")
LINKEDIN = re.compile(r"(?:https?://)?(?:[\w-]+\.)?linkedin\.com/in/[\w%-]+", re.I)
GITHUB = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[\w-]+", re.I)
MAX_CHARS = 20000
MIN_QUOTE = 4  # characters left after normalising; shorter quotes would match anywhere


def extract_text(path):
    """Plain text of a .pdf, .docx or .txt file. PDF and DOCX need the `cv` extra."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise RuntimeError(MISSING.format(".pdf")) from e
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if ext == ".docx":
        try:
            import docx
        except ImportError as e:
            raise RuntimeError(MISSING.format(".docx")) from e
        doc = docx.Document(str(path))
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        return "\n".join([p.text for p in doc.paragraphs] + cells)
    raise ValueError(f"Unsupported file type {ext or path.name!r}; use .pdf, .docx or .txt")


def _norm(text):
    """Lowercase words separated by one space; + and # stay so that C, C++ and C# are different words."""
    return re.sub(r"[^a-z0-9+#]+", " ", str(text).lower()).strip()


def _has(word, norm_text):
    """The (normalised) word is in the normalised text as a whole word."""
    word = _norm(word)
    return bool(word) and re.search(rf"(?<![a-z0-9+#]){re.escape(word)}(?![a-z0-9+#])", norm_text) is not None


def _url(match):
    return match.group() if match.group().startswith("http") else "https://" + match.group()


def _contact(text):
    out = {}
    first = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
    words = first.split()
    if 2 <= len(words) <= 3 and all(w.replace("-", "").isalpha() for w in words):
        out["first_name"], out["last_name"] = words[0], " ".join(words[1:])
    if m := EMAIL.search(text):
        out["email"] = m.group()
    for m in PHONE.finditer(text):
        if sum(c.isdigit() for c in m.group()) >= 9:
            out["phone"] = m.group().strip()
            break
    if m := LINKEDIN.search(text):
        out["linkedin"] = _url(m)
    if m := GITHUB.search(text):
        out["github"] = _url(m)
    return out


MAIN = {"skill": ("name",), "cert": ("name",), "experience": ("title", "org"), "education": ("name",),
        "language": ("name",)}
# Values a quote must contain; one that is not there is dropped from the fact and reported.
CHECKED = {"skill": ("years",), "cert": ("year",), "experience": ("start", "end"), "education": ("org", "year"),
           "language": ("level",)}
MONTHS = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")
PRESENT = ("present", "current", "now", "today", "presente", "actualidad", "actual")
MAX_QUOTE = 300


def _quote(value, haystack):
    """The normalised quote if it is long enough and is in the CV text, else None."""
    q = _norm(value[:MAX_QUOTE]) if isinstance(value, str) else ""
    return q if len(q) >= MIN_QUOTE and q in haystack else None


def _number_in(value, q):
    numbers = {float(n.replace(",", ".")) for n in re.findall(r"\d+(?:[.,]\d+)?", q)}
    return isinstance(value, int | float) and not isinstance(value, bool) and float(value) in numbers


def _date_in(value, q):
    if not isinstance(value, str):
        return False
    if value == "present":
        return any(_has(w, q) for w in PRESENT)
    m = re.fullmatch(r"(\d{4})(?:-(\d{2}))?", value)
    if not m or not _has(m[1], q):
        return False
    return (m[2] is None or f"{m[1]} {m[2]}" in q or f"{m[2]} {m[1]}" in q
            or 1 <= int(m[2]) <= 12 and _has(MONTHS[int(m[2]) - 1], re.sub(r"(?<=[a-z]{3})[a-z]+", "", q)))


def _verified(key, value, q):
    if key == "years":
        return _number_in(value, q)
    if key == "year":
        return isinstance(value, int) and not isinstance(value, bool) and _has(str(value), q)
    if key in ("start", "end"):
        return _date_in(value, q)
    return isinstance(value, str) and _has(value, q)  # org, level


def _fact(raw, haystack, used, log):
    """(fact, reason). The fact is clean, or None with the reason: "unbacked" if its quote is not in the CV or
    does not mention its name, "invalid" if it is malformed. Values the quote does not contain are dropped."""
    kind = raw.get("kind") if isinstance(raw, dict) else None
    if not isinstance(kind, str) or kind not in ID_PREFIX:
        return None, "invalid"
    q = _quote(raw.get("source"), haystack)
    if q is None:
        return None, "unbacked"
    fact = {"kind": kind}
    for key in MAIN[kind]:
        if not isinstance(raw.get(key), str) or not raw[key].strip():
            return None, "invalid"
        if not _has(raw[key], q):
            return None, "unbacked"
        fact[key] = raw[key].strip()
    for key in CHECKED[kind]:
        if key in raw:
            if _verified(key, raw[key], q):
                fact[key] = raw[key].strip() if isinstance(raw[key], str) else raw[key]
            else:
                log(f"unverified: {key} of {fact[MAIN[kind][0]]!r} is not in its quote, left out")
    seed = f"{fact.get('org', '')}-{str(fact.get('start', ''))[:4]}" if kind == "experience" else fact["name"]
    base, n = f"{ID_PREFIX[kind]}.{slug(seed)}", 2
    fact = {"id": base, **fact}
    while fact["id"] in used:
        fact["id"], n = f"{base}-{n}", n + 1
    bullets = []
    for b in raw.get("bullets") if kind == "experience" and isinstance(raw.get("bullets"), list) else []:
        text = " ".join(b["source"][:MAX_QUOTE].split()) if isinstance(b, dict) and isinstance(b.get("source"), str) else ""
        if _quote(text, haystack):
            bullets.append({"id": f"{fact['id']}.b{len(bullets) + 1}", "text": text})
    if bullets:
        fact["bullets"] = bullets
    # An experience whose start was left out is still a draft the interview can finish.
    if fact_errors({"start": "1900-01", **fact} if kind == "experience" else fact, "fact"):
        return None, "invalid"
    used.add(fact["id"])
    return fact, None


def draft_from_text(text, use_llm=True, log=print):
    """A v2 profile draft: contact details by regex and, with `use_llm`, facts the LLM found in the text."""
    draft = {"version": SCHEMA_VERSION, **_contact(text), "facts": []}
    if not use_llm:
        return draft
    reasons = looks_injected(text)
    if reasons:
        log(f"The CV text looks like a prompt injection ({', '.join(reasons)}); skipping the LLM step, "
            "only contact details were read.")
        return draft
    if len(text) > MAX_CHARS:
        log(f"The CV is long: only the first {MAX_CHARS} characters are sent to the LLM (truncated).")
    try:
        items = llm.extract_cv_facts(text[:MAX_CHARS])
    except llm.LLMError as e:
        log(f"Could not ask the LLM for facts ({e}); only contact details were read.")
        return draft
    if not isinstance(items, list):
        log("The LLM reply could not be read as a list of facts; only contact details were read.")
        return draft
    haystack, used, dropped = _norm(text), set(), Counter()
    for item in items:
        fact, reason = _fact(item, haystack, used, log)
        if fact:
            draft["facts"].append(fact)
        else:
            dropped[reason] += 1
    if dropped["unbacked"]:
        log(f"Dropped {dropped['unbacked']} facts not backed by the CV text.")
    if dropped["invalid"]:
        log(f"Dropped {dropped['invalid']} invalid facts.")
    return draft

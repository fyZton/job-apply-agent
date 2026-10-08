import openpyxl
import pytest

from jobagent.tracker import SHEET, Tracker

LINK = "https://www.linkedin.com/jobs/view/1/"


@pytest.mark.parametrize("text", ['=HYPERLINK("http://evil","x")', "+1 dev", "-2+3", "@SUM(A1)", "\t=1", "\r=1"])
def test_formula_like_web_text_is_neutralised(tmp_path, text):
    t = Tracker(tmp_path / "log.xlsx", tmp_path, print)
    t.add_row({"Company": text, "Title": "Dev", "Notes": "= safe? no", "Link": LINK})
    ws = openpyxl.load_workbook(tmp_path / "log.xlsx")[SHEET]
    assert ws.cell(2, 2).data_type == "s" and ws.cell(2, 2).value == "'" + text
    assert ws.cell(2, 3).value == "Dev"  # normal text is untouched
    assert ws.cell(2, 11).value == "'= safe? no"


def test_rows_and_history_survive_a_restart(tmp_path):
    t = Tracker(tmp_path / "log.xlsx", tmp_path, print)
    t.add_row({"Company": "Acme", "Title": "Dev", "Link": LINK})
    t.mark("linkedin:2", "low_fit")
    t.add_sent("linkedin")
    again = Tracker(tmp_path / "log.xlsx", tmp_path, print)
    assert again.seen("linkedin:1") and again.seen("linkedin:2")
    assert again.sent_today("linkedin") == 1
